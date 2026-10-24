"""Author the three-page editable customer description; not part of delivery."""
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / 'artifacts/client_document'
WORK.mkdir(parents=True, exist_ok=True)


def diagram():
    image = Image.new('RGB', (1900, 210), 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 32)
    small = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 25)
    items = [('RGB и время','Кадры видео'),('YOLO26m','Рамки и классы'),('Трекер','Номера объектов'),('Оценка сближения','Рост и движение'),('Индикатор','Области и стрелки')]
    for i, (title, subtitle) in enumerate(items):
        x = 5 + i * 385
        draw.rounded_rectangle((x, 28, x + 345, 180), radius=14, fill='#edf3f5', outline='#516572', width=3)
        for text, y, f in [(title,65,font),(subtitle,113,small)]:
            draw.text((x + (345 - draw.textlength(text,font=f))/2,y),text,font=f,fill='#182b36')
        if i < 4:
            draw.line((x+350,104,x+380,104),fill='#516572',width=5)
            draw.polygon([(x+380,104),(x+368,96),(x+368,112)],fill='#516572')
    p=WORK/'pipeline.png';image.save(p);return p


doc = Document()
section = doc.sections[0]
section.page_width,section.page_height=Mm(210),Mm(297)
section.top_margin,section.bottom_margin=Mm(16),Mm(16)
section.left_margin,section.right_margin=Mm(18),Mm(18)
section.header_distance,section.footer_distance=Mm(7),Mm(8)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.line_spacing=1.04
for name,size in [('Title',20),('Heading 1',13),('Heading 2',11)]:
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].paragraph_format.space_before=Pt(9 if name!='Title' else 0)
    doc.styles[name].paragraph_format.space_after=Pt(5)
doc.styles['Caption'].font.size=Pt(8.5)
footer=section.footer.paragraphs[0];footer.alignment=2
footer.add_run('AirAntiCollision   •   ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for r in footer.runs:r.font.size=Pt(8)


def p(text='',bold=False,size=None):
    row=doc.add_paragraph();r=row.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    return row


def h(text):doc.add_heading(text,1)


def link(paragraph,text,url):
    rel=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    node=OxmlElement('w:hyperlink');node.set(qn('r:id'),rel)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'175675');props.append(color)
    run.append(props);t=OxmlElement('w:t');t.text=text;run.append(t);node.append(run);paragraph._p.append(node)


def code(files):
    row=doc.add_paragraph();row.paragraph_format.space_after=Pt(5)
    r=row.add_run('Код: ');r.font.size=Pt(8.5)
    for i,(name,symbol) in enumerate(files):
        if i:row.add_run('; ').font.size=Pt(8.5)
        r=row.add_run(name+' · '+symbol)
        r.font.name='Consolas';r.font.size=Pt(9)
    for r in row.runs:r.font.size=Pt(8.5)


def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Mm(width)
    for cell,text in zip(t.rows[0].cells,headers):cell.text=text
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=text
    for index,row in enumerate(t.rows):
        for cell,width in zip(row.cells,widths):
            cell.width=Mm(width)
            tcPr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                item=OxmlElement('w:'+side);item.set(qn('w:w'),'70');item.set(qn('w:type'),'dxa');margins.append(item)
            tcPr.append(margins)
            if index==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EFF2');tcPr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                para.paragraph_format.line_spacing=1.0
                for run in para.runs:run.font.size=Pt(9);run.bold=index==0
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t


doc.add_paragraph('Визуальная оценка сближения самолётов','Title')
p('AirAntiCollision  •  Описание алгоритма и запуска  •  24 октября 2026',size=9)
p('Система обрабатывает записанное RGB-видео: находит воздушные объекты, сохраняет их номера между кадрами и выделяет признаки возможного сближения. Результат - видео с индикатором и журнал наблюдений. Это визуальная оценка ситуации; физическая траектория столкновения по одному некалиброванному видео не устанавливается.')
doc.add_picture(str(diagram()),width=Mm(174))
h('Детектор и выбранная архитектура')
table(['Режим','Назначение','Фактический вход NCHW'],[
    ['B · YOLO26m-P2','Основной для NPS: самолётные аппараты; в интерфейсе «самолёт».','1 × 3 × 800 × 1440'],
    ['A · YOLO26m','Альтернатива: самолёт, вертолёт, птица, аэростат, парашют.','1 × 3 × 416 × 704'],
],[37,84,53])
p('Оба экспортированных детектора [6] работают через TensorRT без дообучения. B обрабатывает кадр целиком; A дополнительно просматривает перекрывающиеся фрагменты исходного разрешения. Координаты рамок возвращаются в исходный кадр.')
code([('detector_p2.py','DetectorP2'),('detector.py','Detector')])
p('Для NPS выбрана связка P2 с собственным трекером. На трёх оценочных клипах 020/034/043 она дала IDF1 82,59% и пять смен ID. Это результат сохранения идентичности объектов в выбранной выборке, а не точность определения столкновений.')
h('Связывание наблюдений и контроль камеры')
p('Трекер сравнивает новую рамку с прогнозом центра по последним наблюдениям. Учитываются расстояние, изменение размера и класс; подходящие пары назначаются по возрастанию стоимости. При неоднозначном соответствии история геометрии сбрасывается. Новые, пропавшие и обрезанные границей объекты не получают достоверную оценку роста.')
p('В режиме auto движение фона проверяется по оптическому потоку KLT и преобразованию подобия с RANSAC. Ненадёжная оценка камеры сбрасывает историю. Такая проверка не отделяет вращение самолёта от переноса камеры и не заменяет IMU. Режим fixed применён только к неподвижной камере QUT.')
code([('tracking.py','Tracker.update / CameraMotion.update')])

doc.add_page_break()
h('Как формируется признак сближения')
p('Для каждого ID используются только свежие рамки детектора и исходные временные метки. За окно до 0,6 с оцениваются рост логарифма большего размера рамки и скорость её центра. Наклон берётся как медиана попарных оценок, чтобы уменьшить влияние дрожания. Анализ начинается после 0,3 с и не менее пяти измерений; минимальный размер - 6 пикселей.')
p('Кандидат на сближение должен устойчиво расти не медленнее 0,035 в секунду, иметь небольшое смещение относительно своего размера и стабильную форму рамки. Рост должен превышать оценку шума и подтверждаться ещё 0,16 с. Малое смещение измеряется в изображении, без перевода в физический пеленг. Правило разработано для демонстрации на основе идей оптического расширения [4, 5].')
table(['Индикация','Что наблюдалось'],[
    ['КОНТРОЛЬ','Объект наблюдается устойчиво; подтверждённого признака сближения нет.'],
    ['ВОЗМ. СБЛИЖЕНИЕ','Устойчивый рост при малом относительном смещении изображения.'],
    ['НЕТ ДАННЫХ','История, геометрия, ID или оценка камеры недостаточно надёжны.'],
],[47,127])
code([('demo_cue.py','EarlyEncounterCue.update')])
doc.add_picture(str(ROOT/'demo/aviation_nps_field/stable_tracking.png'),width=Mm(174))
doc.add_paragraph('NPS 032. Бортовая камера: номера, реальные ROI и направление разделения в центральном поле.', 'Caption')
p('Приоритетный объект выбирается с временным подтверждением. Крупная стрелка направлена от его прогнозируемого положения в кадре через 0,5 с; противоречивые или ненадёжные данные скрывают её. Стрелка может отображаться и в состоянии «КОНТРОЛЬ». Она показывает разделение в изображении, а не рассчитанный безопасный манёвр.')
doc.add_picture(str(ROOT/'demo/aviation_approach_field/encounter_cue.png'),width=Mm(174))
doc.add_paragraph('QUT M2. Измеренный рост объекта даёт янтарную индикацию; камера в этой записи неподвижна.', 'Caption')
p('Перекрывающиеся области рамок с коротким прогнозом движения объединяются в связные группы и обводятся общим контуром. Янтарный цвет означает признак сближения, голубой - наблюдение. Пунктир показывает движение объектов. Неподсвеченная область не считается проверенной безопасной зоной.',size=9.5)
code([('aviation_regions.py','compute_attention_regions'),('aviation_render.py','AviationRenderer')])

doc.add_page_break()
h('Состав комплекта и запуск')
p('Для просмотра достаточно открыть Watch_NPS.cmd или Watch_QUT.cmd. Готовые ролики находятся в demo; устанавливать Python для их просмотра не нужно. Файлы примеров и модели включены в комплект.')
table(['Папка или файл','Содержимое'],[
    ['app','Исходный код обработки; run.py - точка входа; prepare_models.py - сборка TensorRT.'],
    ['models','Два ONNX, два engine и реестр контрольных сумм.'],
    ['examples / demo','Исходные примеры NPS и QUT / готовые видео и кадры индикатора.'],
    ['results','Новые MP4, JSONL наблюдений и решений, параметры и статус запуска.'],
],[43,131])
p('Для повторной обработки нужны Windows x64, Python 3.12, NVIDIA с драйвером для CUDA 13 и FFmpeg/FFprobe в PATH. Setup.cmd создаёт отдельную среду и пересобирает TensorRT-модели для установленной видеокарты. FFmpeg устанавливается отдельно: winget install --id Gyan.FFmpeg --exact. После настройки запускаются Run_NPS.cmd или Run_QUT.cmd.')
p('Свой файл: python app/run.py --video C:/video/flight.mp4 --model B. Для пяти классов выбирается --model A. Сборка проверена в рабочей среде Windows / RTX 4060 / TensorRT 11.4; установка на чистом ПК отдельно не проверялась. Предобработка поставки выполняется на CPU, нейросетевой вывод - на GPU.',size=9.5)
h('Как понимать результат')
p('Размер рамки зависит не только от сближения, но и от ракурса, размытия и ошибок детектора. Без телеметрии и калибровки не вычисляются доказанная 3D-траектория, вероятность столкновения или безопасная команда управления. Автопилот и модель динамики самолёта не подключены.')
p('Только для NPS на индикатор выводится примерная дальность: приняты размах 1,5 м и приближённая горизонтальная эквидистантная проекция с FOV 120° для GoPro HERO3 Ultra Wide. Это не калибровка конкретной камеры. Оценка сглаживается и не участвует в предупреждении или выборе стрелки. Код: app/airavoid/demo_range.py.')
h('Данные и научные источники')
p('NPS-Drones [1] с уточнёнными рамками Dogfight [2] и поклиповыми ID использовался для оценки сопровождения. QUT [3] использовался как пример наблюдаемого сближения с неподвижной камерой. Эталонные рамки и ID в обычный конвейер не передаются. Представленные записи не доказывают физическую точность классификации столкновений.',size=9.5)
sources=[
    ('[1] Li et al. Fast and Robust UAV to UAV Detection and Tracking from Video, TETC, 2022. ', 'NPS / Purdue dataset', 'https://engineering.purdue.edu/~bouman/UAV_Dataset/'),
    ('[2] Ashraf, Sultani, Shah. Dogfight: Detecting Drones From Drones Videos, CVPR, 2021. ', 'Статья и уточнение разметки', 'https://openaccess.thecvf.com/content/CVPR2021/papers/Ashraf_Dogfight_Detecting_Drones_From_Drones_Videos_CVPR_2021_paper.pdf'),
    ('[3] Martin, Riseley, Ford. A Dataset of Stationary, Fixed-wing Aircraft on a Collision Course, ICUAS, 2022. ', 'QUT dataset', 'https://open.qcr.ai/dataset/aircraft-collision-course/'),
    ('[4] Bauer et al. Monocular Image-based Time to Collision and Closest Point of Approach Estimation, MED, 2016. ', 'DOI 10.1109/MED.2016.7535965', 'https://doi.org/10.1109/MED.2016.7535965'),
    ('[5] Ghosh et al. AirTrack: Onboard Deep Learning Framework for Long-Range Aircraft Detection and Tracking, 2023. ', 'arXiv:2209.12849', 'https://arxiv.org/abs/2209.12849'),
    ('[6] Ultralytics. YOLO26. ', 'Описание семейства детекторов', 'https://docs.ultralytics.com/models/yolo26/'),
]
for text,label,url in sources:
    row=p(text,size=8.5);row.paragraph_format.space_after=Pt(3);link(row,label,url)
p('Условия сторонних моделей и данных сохранены в licenses. Экспорт YOLO содержит указание AGPL-3.0; NPS и QUT опубликованы под BSD-3-Clause.',size=8)

target=WORK/'Algorithm_RU.docx';doc.save(target);print(target)
